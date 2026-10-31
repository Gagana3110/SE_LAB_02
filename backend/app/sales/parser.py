"""Reads an uploaded sales file into rows, refusing anything that is not a genuine CSV or XLSX."""

import csv
import io
import zipfile
from collections.abc import Iterator
from datetime import date, datetime
from pathlib import PurePath

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_UNZIPPED_BYTES = 200 * 1024 * 1024  # guards against a small file that expands enormously
REQUIRED_COLUMNS = ("date", "product_code", "quantity")
MAX_QUANTITY = 2_147_483_647
MAX_REGION_LENGTH = 60


class UnsupportedFile(Exception):
    """Not a CSV or XLSX file (HTTP 415)."""


class BadLayout(Exception):
    """A readable file with missing columns or no data (HTTP 400)."""


def read_rows(filename: str, content: bytes) -> Iterator[tuple[int, dict]]:
    """Yields (row number as seen in a spreadsheet, {column: value})."""
    extension = PurePath(filename or "").suffix.lower()
    if extension == ".csv":
        rows = _read_csv(content)
    elif extension == ".xlsx":
        rows = _read_xlsx(content)
    else:
        raise UnsupportedFile("Only .csv and .xlsx files are accepted")
    return _with_columns(rows)


def _read_csv(content: bytes) -> list[tuple[int, list]]:
    if b"\0" in content:
        raise UnsupportedFile("This file is not a valid CSV file")
    try:
        text = content.decode("utf-8-sig")
        reader = csv.reader(io.StringIO(text, newline=""))
        return [(reader.line_num, row) for row in reader]
    except (UnicodeDecodeError, csv.Error) as exc:
        raise UnsupportedFile("This file is not a valid UTF-8 CSV file") from exc


def _read_xlsx(content: bytes) -> list[tuple[int, list]]:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
    except zipfile.BadZipFile as exc:
        raise UnsupportedFile("This file is not a valid Excel workbook") from exc

    names = {entry.filename.lower() for entry in entries}
    if "xl/workbook.xml" not in names:
        raise UnsupportedFile("This file is not a valid Excel workbook")
    if "xl/vbaproject.bin" in names:
        raise UnsupportedFile("Workbooks containing macros are not accepted")
    if sum(entry.file_size for entry in entries) > MAX_UNZIPPED_BYTES:
        raise UnsupportedFile("This workbook is too large when unpacked")

    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        try:
            sheet = workbook.worksheets[0]
            return [(number, list(row)) for number, row in
                    enumerate(sheet.iter_rows(values_only=True), start=1)]
        finally:
            workbook.close()
    except (InvalidFileException, zipfile.BadZipFile, IndexError, KeyError, ValueError,
            OSError, TypeError, AttributeError) as exc:
        raise UnsupportedFile("This file is not a valid Excel workbook") from exc


def _is_blank(row: list) -> bool:
    return all(cell is None or str(cell).strip() == "" for cell in row)


def _with_columns(rows: list[tuple[int, list]]) -> Iterator[tuple[int, dict]]:
    filled = [(number, row) for number, row in rows if not _is_blank(row)]
    if not filled:
        raise BadLayout("The file is empty")

    header = [str(cell or "").strip().lower() for cell in filled[0][1]]
    missing = [column for column in REQUIRED_COLUMNS if column not in header]
    if missing:
        raise BadLayout(f"Missing columns: {', '.join(missing)}")
    if len(filled) == 1:
        raise BadLayout("The file has no data rows")
    return ((number, dict(zip(header, row))) for number, row in filled[1:])


def parse_date(value) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return date.fromisoformat(str(value or "").strip())


def parse_quantity(value) -> int:
    text = str(value if value is not None else "").strip()
    number = float(text)  # also accepts the 12.0 that Excel produces for whole numbers
    if not number.is_integer():
        raise ValueError(text)
    return int(number)
